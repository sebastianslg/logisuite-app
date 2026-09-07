"""
data_tools.py
Herramientas transversales de datos:
  - Búsqueda global across todas las tablas del sistema.
  - Importación masiva desde CSV/Excel con validación fila por fila.
  - Exportación consolidada a un único Excel multi-hoja.
"""
import io
from datetime import datetime
from typing import List, Tuple

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from database.db import run_query, run_write


# ---------------------------------------------------------------------------
# BÚSQUEDA GLOBAL
# ---------------------------------------------------------------------------
def global_search(term: str) -> List[dict]:
    """Busca un término en todas las entidades relevantes del sistema.

    Devuelve una lista unificada de resultados con el módulo, tipo de entidad
    y una descripción legible, para que el usuario ubique cualquier registro
    sin saber en qué módulo está."""
    if not term or len(term.strip()) < 2:
        return []
    like = f"%{term.strip()}%"
    resultados = []

    # Inventario (SKU y nombre de producto)
    df = run_query("""
        SELECT i.item_id, i.sku, i.name, i.quantity, n.name AS almacen
        FROM inventory_items i
        JOIN warehouses w ON i.warehouse_id = w.warehouse_id
        JOIN nodes n ON w.node_id = n.node_id
        WHERE i.sku LIKE ? OR i.name LIKE ?
    """, (like, like))
    for r in df.to_dict(orient="records"):
        resultados.append({"modulo": "Warehouse", "tipo": "Ítem de inventario",
                            "identificador": r["sku"],
                            "descripcion": f"{r['name']} — {r['quantity']:.0f} u. en {r['almacen']}",
                            "id": r["item_id"]})

    # Vehículos (placa y tipo)
    df = run_query("SELECT * FROM vehicles WHERE plate LIKE ? OR vehicle_type LIKE ?", (like, like))
    for r in df.to_dict(orient="records"):
        resultados.append({"modulo": "Fleet", "tipo": "Vehículo", "identificador": r["plate"],
                            "descripcion": f"{r['vehicle_type']} — {r['status']} — {r['odometer_km']:.0f} km",
                            "id": r["vehicle_id"]})

    # Conductores
    df = run_query("SELECT * FROM drivers WHERE name LIKE ? OR license_number LIKE ?", (like, like))
    for r in df.to_dict(orient="records"):
        resultados.append({"modulo": "Fleet", "tipo": "Conductor", "identificador": r["license_number"],
                            "descripcion": f"{r['name']} — categoría {r['license_category']} — vence {r['license_expiry']}",
                            "id": r["driver_id"]})

    # Documentos aduaneros
    df = run_query("SELECT * FROM customs_documents WHERE doc_number LIKE ? OR doc_type LIKE ?",
                    (like, like))
    for r in df.to_dict(orient="records"):
        resultados.append({"modulo": "Customs", "tipo": "Documento aduanero",
                            "identificador": r["doc_number"],
                            "descripcion": f"{r['doc_type']} — {r['status']} — envío #{r['shipment_id']}",
                            "id": r["doc_id"]})

    # Nodos de la red
    df = run_query("SELECT * FROM nodes WHERE name LIKE ? OR city LIKE ? OR node_type LIKE ?",
                    (like, like, like))
    for r in df.to_dict(orient="records"):
        resultados.append({"modulo": "Red", "tipo": f"Nodo ({r['node_type']})",
                            "identificador": r["name"],
                            "descripcion": f"{r['city']} — {'activo' if r['active'] else 'inactivo'}",
                            "id": r["node_id"]})

    # Envíos (por ID exacto si el término es numérico)
    if term.strip().isdigit():
        df = run_query("""
            SELECT s.*, no.name AS origen, nd.name AS destino FROM shipments s
            JOIN nodes no ON s.origin_node_id = no.node_id
            JOIN nodes nd ON s.dest_node_id = nd.node_id
            WHERE s.shipment_id = ?
        """, (int(term.strip()),))
        for r in df.to_dict(orient="records"):
            resultados.append({"modulo": "Freight", "tipo": "Envío",
                                "identificador": f"#{r['shipment_id']}",
                                "descripcion": f"{r['origen']} → {r['destino']} — {r['status']}",
                                "id": r["shipment_id"]})

    return resultados


# ---------------------------------------------------------------------------
# IMPORTACIÓN MASIVA CON VALIDACIÓN
# ---------------------------------------------------------------------------
IMPORT_SCHEMAS = {
    "inventario": {
        "columnas_requeridas": ["sku", "name", "warehouse_id", "zone", "quantity",
                                  "min_stock", "max_stock", "unit_cost"],
        "tabla": "inventory_items",
        "descripcion": "Ítems de inventario. La zona debe ser Alta, Media o Baja.",
    },
    "vehiculos": {
        "columnas_requeridas": ["plate", "vehicle_type", "capacity_kg", "capacity_m3"],
        "tabla": "vehicles",
        "descripcion": "Vehículos de la flota. El tipo debe ser una tipología válida.",
    },
    "conductores": {
        "columnas_requeridas": ["name", "license_number", "license_category", "license_expiry"],
        "tabla": "drivers",
        "descripcion": "Conductores. La vigencia debe tener formato AAAA-MM-DD.",
    },
    "nodos": {
        "columnas_requeridas": ["name", "node_type", "latitude", "longitude"],
        "tabla": "nodes",
        "descripcion": "Nodos de la red. El tipo debe ser Planta, Almacen, CD, Cliente, Hub o Gateway.",
    },
}

ZONAS_VALIDAS = ("Alta", "Media", "Baja")
TIPOS_VEHICULO = ("Camion 2 ejes", "Camion 3 ejes", "Tractomula", "Furgon", "Van")
TIPOS_NODO = ("Planta", "Almacen", "CD", "Cliente", "Hub", "Gateway")


def validate_import(df: pd.DataFrame, tipo: str) -> Tuple[List[dict], List[dict]]:
    """Valida un DataFrame antes de importarlo.

    Devuelve (filas_validas, errores). Cada error indica el número de fila y el
    motivo, para que el usuario corrija su archivo sin adivinar."""
    if tipo not in IMPORT_SCHEMAS:
        return [], [{"fila": 0, "error": f"Tipo de importación desconocido: {tipo}"}]

    esquema = IMPORT_SCHEMAS[tipo]
    requeridas = esquema["columnas_requeridas"]

    faltantes = [c for c in requeridas if c not in df.columns]
    if faltantes:
        return [], [{"fila": 0, "error": f"Faltan columnas obligatorias: {', '.join(faltantes)}"}]

    validas, errores = [], []
    for idx, row in df.iterrows():
        fila_num = idx + 2  # +2: fila 1 es el encabezado en el archivo original
        fila_errores = []

        # Campos vacíos
        for col in requeridas:
            if pd.isna(row[col]) or str(row[col]).strip() == "":
                fila_errores.append(f"'{col}' está vacío")

        if not fila_errores:
            # Validaciones específicas por tipo
            if tipo == "inventario":
                if str(row["zone"]) not in ZONAS_VALIDAS:
                    fila_errores.append(f"zona '{row['zone']}' inválida (use: {', '.join(ZONAS_VALIDAS)})")
                for col in ("quantity", "min_stock", "max_stock", "unit_cost"):
                    try:
                        if float(row[col]) < 0:
                            fila_errores.append(f"'{col}' no puede ser negativo")
                    except (ValueError, TypeError):
                        fila_errores.append(f"'{col}' debe ser numérico")
                try:
                    if float(row["min_stock"]) > float(row["max_stock"]):
                        fila_errores.append("min_stock no puede ser mayor que max_stock")
                except (ValueError, TypeError):
                    pass
                wh = run_query("SELECT warehouse_id FROM warehouses WHERE warehouse_id = ?",
                                (int(row["warehouse_id"]),)) if str(row["warehouse_id"]).strip().isdigit() else pd.DataFrame()
                if wh.empty:
                    fila_errores.append(f"warehouse_id {row['warehouse_id']} no existe")

            elif tipo == "vehiculos":
                if str(row["vehicle_type"]) not in TIPOS_VEHICULO:
                    fila_errores.append(f"tipo '{row['vehicle_type']}' inválido")
                existing = run_query("SELECT vehicle_id FROM vehicles WHERE plate = ?",
                                      (str(row["plate"]),))
                if not existing.empty:
                    fila_errores.append(f"la placa {row['plate']} ya existe")
                for col in ("capacity_kg", "capacity_m3"):
                    try:
                        if float(row[col]) <= 0:
                            fila_errores.append(f"'{col}' debe ser mayor que cero")
                    except (ValueError, TypeError):
                        fila_errores.append(f"'{col}' debe ser numérico")

            elif tipo == "conductores":
                try:
                    datetime.strptime(str(row["license_expiry"]).split(" ")[0], "%Y-%m-%d")
                except ValueError:
                    fila_errores.append("license_expiry debe tener formato AAAA-MM-DD")
                existing = run_query("SELECT driver_id FROM drivers WHERE license_number = ?",
                                      (str(row["license_number"]),))
                if not existing.empty:
                    fila_errores.append(f"la licencia {row['license_number']} ya existe")

            elif tipo == "nodos":
                if str(row["node_type"]) not in TIPOS_NODO:
                    fila_errores.append(f"tipo '{row['node_type']}' inválido")
                try:
                    lat, lon = float(row["latitude"]), float(row["longitude"])
                    if not (-90 <= lat <= 90):
                        fila_errores.append("latitud fuera de rango (-90 a 90)")
                    if not (-180 <= lon <= 180):
                        fila_errores.append("longitud fuera de rango (-180 a 180)")
                except (ValueError, TypeError):
                    fila_errores.append("latitud/longitud deben ser numéricas")

        if fila_errores:
            errores.append({"fila": fila_num, "error": "; ".join(fila_errores),
                             "datos": str(dict(row))[:120]})
        else:
            validas.append(row.to_dict())

    return validas, errores


def execute_import(filas: List[dict], tipo: str) -> int:
    """Inserta las filas ya validadas. Devuelve el número de registros creados."""
    insertados = 0
    for row in filas:
        try:
            if tipo == "inventario":
                run_write("""INSERT INTO inventory_items (warehouse_id, sku, name, zone, quantity,
                             min_stock, max_stock, unit_cost) VALUES (?,?,?,?,?,?,?,?)""",
                           (int(row["warehouse_id"]), str(row["sku"]), str(row["name"]),
                            str(row["zone"]), float(row["quantity"]), float(row["min_stock"]),
                            float(row["max_stock"]), float(row["unit_cost"])))
            elif tipo == "vehiculos":
                run_write("""INSERT INTO vehicles (plate, vehicle_type, capacity_kg, capacity_m3,
                             status, odometer_km) VALUES (?,?,?,?,'Disponible',?)""",
                           (str(row["plate"]), str(row["vehicle_type"]), float(row["capacity_kg"]),
                            float(row["capacity_m3"]), float(row.get("odometer_km", 0) or 0)))
            elif tipo == "conductores":
                run_write("""INSERT INTO drivers (name, license_number, license_category,
                             license_expiry, status) VALUES (?,?,?,?,'Activo')""",
                           (str(row["name"]), str(row["license_number"]),
                            str(row["license_category"]),
                            str(row["license_expiry"]).split(" ")[0]))
            elif tipo == "nodos":
                run_write("""INSERT INTO nodes (name, node_type, latitude, longitude, city, active)
                             VALUES (?,?,?,?,?,1)""",
                           (str(row["name"]), str(row["node_type"]), float(row["latitude"]),
                            float(row["longitude"]), str(row.get("city", "") or "")))
            insertados += 1
        except Exception:
            continue
    return insertados


def import_template(tipo: str) -> bytes:
    """Genera una plantilla Excel vacía con las columnas correctas."""
    esquema = IMPORT_SCHEMAS.get(tipo, {})
    cols = esquema.get("columnas_requeridas", [])
    df = pd.DataFrame(columns=cols)
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Plantilla")
    return buf.getvalue()


# ---------------------------------------------------------------------------
# EXPORTACIÓN CONSOLIDADA MULTI-HOJA
# ---------------------------------------------------------------------------
EXPORT_QUERIES = {
    "Nodos": "SELECT * FROM nodes",
    "Corredores": """SELECT c.corridor_id, no.name AS origen, nd.name AS destino,
                      c.distance_km, c.transit_time_h, c.mode, c.cost_per_km, c.active
                      FROM corridors c JOIN nodes no ON c.origin_node_id=no.node_id
                      JOIN nodes nd ON c.dest_node_id=nd.node_id""",
    "Inventario": """SELECT i.sku, i.name, n.name AS almacen, i.zone, i.quantity,
                      i.min_stock, i.max_stock, i.unit_cost,
                      i.quantity*i.unit_cost AS valoracion
                      FROM inventory_items i
                      JOIN warehouses w ON i.warehouse_id=w.warehouse_id
                      JOIN nodes n ON w.node_id=n.node_id""",
    "Movimientos": """SELECT m.movement_date, i.sku, i.name, m.movement_type, m.quantity, m.reference
                       FROM warehouse_movements m JOIN inventory_items i ON m.item_id=i.item_id
                       ORDER BY m.movement_date DESC""",
    "Envios": """SELECT s.shipment_id, no.name AS origen, nd.name AS destino, s.cargo_units,
                  s.weight_kg, s.volume_m3, s.status, s.promised_date, s.delivered_date,
                  s.transaction_cost, s.distance_friction_cost, s.shipment_cost, s.total_cost
                  FROM shipments s JOIN nodes no ON s.origin_node_id=no.node_id
                  JOIN nodes nd ON s.dest_node_id=nd.node_id""",
    "Rutas": """SELECT r.route_id, r.shipment_id, r.algorithm, r.total_distance_km, r.total_cost,
                 r.eta_hours, r.actual_hours, v.plate, d.name AS conductor
                 FROM routes r LEFT JOIN vehicles v ON r.vehicle_id=v.vehicle_id
                 LEFT JOIN drivers d ON r.driver_id=d.driver_id""",
    "Vehiculos": "SELECT * FROM vehicles",
    "Conductores": "SELECT * FROM drivers",
    "Mantenimientos": """SELECT m.maintenance_date, v.plate, m.maintenance_type, m.odometer_km,
                          m.cost, m.description FROM maintenance_records m
                          JOIN vehicles v ON m.vehicle_id=v.vehicle_id""",
    "Doc Aduaneros": "SELECT * FROM customs_documents",
    "Aranceles": "SELECT * FROM customs_duties",
}


def export_all_to_excel() -> bytes:
    """Genera un único Excel con una hoja por módulo, con formato profesional."""
    wb = Workbook()
    wb.remove(wb.active)

    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for sheet_name, sql in EXPORT_QUERIES.items():
        try:
            df = run_query(sql)
        except Exception:
            continue
        ws = wb.create_sheet(title=sheet_name[:31])

        if df.empty:
            ws.cell(row=1, column=1, value="Sin registros")
            continue

        for j, col in enumerate(df.columns, start=1):
            c = ws.cell(row=1, column=j, value=str(col))
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="2A9D8F")
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = border

        for i, (_, row) in enumerate(df.iterrows(), start=2):
            for j, val in enumerate(row, start=1):
                c = ws.cell(row=i, column=j, value=val)
                c.border = border
                c.alignment = Alignment(horizontal="center")
                if i % 2 == 0:
                    c.fill = PatternFill("solid", fgColor="F2F2F2")

        for j, col in enumerate(df.columns, start=1):
            max_len = max([len(str(col))] + [len(str(v)) for v in df[col].astype(str)])
            ws.column_dimensions[get_column_letter(j)].width = min(max(max_len + 3, 10), 38)
        ws.freeze_panes = ws.cell(row=2, column=1)

    if not wb.sheetnames:
        wb.create_sheet(title="Vacio")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
