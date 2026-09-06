"""
report_exporter.py
Generación de reportes exportables en Excel (openpyxl, con formato profesional:
encabezados de color, bordes, autoajuste de columnas) y PDF (reportlab) para
cualquier módulo del sistema.
"""
import io
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm

HEADER_FILL = "2A9D8F"
HEADER_FONT_COLOR = "FFFFFF"


def dataframe_to_excel_bytes(df: pd.DataFrame, sheet_name: str = "Reporte", title: str = None) -> bytes:
    """Convierte un DataFrame en un archivo Excel con estilo académico/ejecutivo."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]

    start_row = 1
    if title:
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(df.columns), 1))
        cell = ws.cell(row=1, column=1, value=title)
        cell.font = Font(size=14, bold=True, color=HEADER_FONT_COLOR)
        cell.fill = PatternFill("solid", fgColor="264653")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 26
        start_row = 3

    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for j, col_name in enumerate(df.columns, start=1):
        c = ws.cell(row=start_row, column=j, value=str(col_name))
        c.font = Font(bold=True, color=HEADER_FONT_COLOR)
        c.fill = PatternFill("solid", fgColor=HEADER_FILL)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border

    for i, (_, row) in enumerate(df.iterrows(), start=start_row + 1):
        for j, val in enumerate(row, start=1):
            c = ws.cell(row=i, column=j, value=val)
            c.border = border
            c.alignment = Alignment(horizontal="center")
            if i % 2 == 0:
                c.fill = PatternFill("solid", fgColor="F2F2F2")

    for j, col_name in enumerate(df.columns, start=1):
        max_len = max([len(str(col_name))] + [len(str(v)) for v in df[col_name].astype(str)])
        ws.column_dimensions[get_column_letter(j)].width = min(max(max_len + 3, 10), 40)

    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def dataframe_to_pdf_bytes(df: pd.DataFrame, title: str = "Reporte") -> bytes:
    """Convierte un DataFrame en un PDF con tabla estilizada, orientación horizontal."""
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(letter),
                             topMargin=1.5 * cm, bottomMargin=1.5 * cm)
    styles = getSampleStyleSheet()
    elements = [Paragraph(title, styles["Title"]), Spacer(1, 12)]

    data = [list(df.columns)] + df.astype(str).values.tolist()
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2A9D8F")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7.5),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#BFBFBF")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F2F2")]),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    elements.append(table)
    doc.build(elements)
    return buf.getvalue()
