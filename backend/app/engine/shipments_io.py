"""
shipments_io.py
Carga masiva de envíos desde Excel (.xlsx) o CSV.

Columnas (encabezados en minúscula, sin importar el orden):
  obligatorias: origen, destino, carga, cliente, peso_t, salida
  opcionales:   prioridad (tiempo | costo | balanceado), retraso_h, modos, vias

Primero se valida todo sin guardar (vista previa con la ruta que tendría cada fila); al confirmar,
cada fila válida se crea con `operations.create_shipment`, igual que desde el formulario, así que
cada envío queda con su ruta, su código y su línea en el historial. Las filas con error se
informan y no detienen las demás.
"""
from __future__ import annotations

import base64
import binascii
import io
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.api import format as F
from app.engine import operations as ops
from app.engine.common import OperationError

REQUIRED = {"origen", "destino", "carga", "cliente", "peso_t", "salida"}
OPTIONAL = ["prioridad", "retraso_h", "modos", "vias"]
MAX_ROWS = 500
ALIASES = {"peso": "peso_t", "peso (t)": "peso_t", "fecha_salida": "salida", "salida_at": "salida",
           "retraso": "retraso_h", "vías": "vias"}


def _text(v) -> str:
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return ""
    return str(v).strip()


def _departure(v) -> str:
    if isinstance(v, (pd.Timestamp, datetime)):
        return pd.Timestamp(v).to_pydatetime().isoformat(timespec="minutes")
    s = _text(v)
    if not s:
        raise OperationError("Falta la fecha de salida.")
    for fmt in ("%Y-%m-%d %H:%M", "%Y-%m-%dT%H:%M", "%Y-%m-%d", "%d/%m/%Y %H:%M", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, fmt).isoformat(timespec="minutes")
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(s).isoformat(timespec="minutes")
    except ValueError:
        raise OperationError(f"Fecha de salida no reconocida: '{s}'. Usa AAAA-MM-DD HH:MM o DD/MM/AAAA HH:MM.")


def _number(v, field: str, default: float | None = None) -> float:
    s = _text(v)
    if not s:
        if default is None:
            raise OperationError(f"Falta {field}.")
        return default
    try:
        return float(s.replace(",", "."))
    except ValueError:
        raise OperationError(f"{field} debe ser un número ('{s}').")


def _read(filename: str, content_b64: str) -> pd.DataFrame:
    try:
        raw = base64.b64decode(content_b64, validate=True)
    except (binascii.Error, ValueError):
        raise OperationError("El archivo llegó dañado (base64 inválido).")
    try:
        if filename.lower().endswith((".xlsx", ".xlsm")):
            df = pd.read_excel(io.BytesIO(raw), dtype=object)
        else:
            df = pd.read_csv(io.BytesIO(raw), dtype=object, sep=None, engine="python")
    except Exception as exc:  # archivo ilegible: se informa sin tumbar la API
        raise OperationError(f"No se pudo leer el archivo: {exc}")
    df.columns = [ALIASES.get(str(c).strip().lower(), str(c).strip().lower()) for c in df.columns]
    missing = REQUIRED - set(df.columns)
    if missing:
        raise OperationError(f"Faltan columnas: {', '.join(sorted(missing))}.")
    df = df.dropna(how="all")
    if len(df) > MAX_ROWS:
        raise OperationError(f"El archivo tiene {len(df)} filas; el máximo por carga es {MAX_ROWS}.")
    if df.empty:
        raise OperationError("El archivo no tiene filas.")
    return df


def _row_data(row) -> dict:
    return {
        "origin": _text(row.get("origen")), "destination": _text(row.get("destino")),
        "cargo": _text(row.get("carga")), "client": _text(row.get("cliente")),
        "weight_t": _number(row.get("peso_t"), "peso_t"),
        "priority": (_text(row.get("prioridad")).lower() or "balanceado"),
        "departure_at": _departure(row.get("salida")),
        "delay_h": _number(row.get("retraso_h"), "retraso_h", 0.0),
        "forced_modes": [m.strip().lower() for m in _text(row.get("modos")).split(",") if m.strip()],
        "forced_corridors": [c.strip() for c in _text(row.get("vias")).split(";") if c.strip()],
    }


def import_shipments(net, filename: str, content_b64: str, user: str | None, now: datetime, commit: bool) -> dict:
    df = _read(filename, content_b64)
    results, created = [], []
    for i, (_, row) in enumerate(df.iterrows()):
        line = i + 2  # fila en la hoja (encabezado en la 1)
        try:
            data = _row_data(row)
            route = ops.preview(net, data)
            item = {"row": line, "ok": True, "origin": data["origin"], "destination": data["destination"],
                    "cargo": data["cargo"], "client": data["client"], "weight_t": data["weight_t"],
                    "departure_at": data["departure_at"],
                    "modes": [lg["mode"] for lg in route["legs"]],
                    "time_fmt": F.duration_fmt(route["time_h"]), "cost": F.money(route["cost"])}
            if commit:
                code = ops.create_shipment(net, data, user, now)
                item["code"] = code
                created.append(code)
            results.append(item)
        except OperationError as exc:
            results.append({"row": line, "ok": False, "error": str(exc)})
    valid = sum(1 for r in results if r["ok"])
    if commit and created:
        ops.audit(user, "IMPORT", "mm_shipments", None,
                  f"Carga masiva desde {filename}: {len(created)} envío(s) creados, {len(results) - valid} con error")
    return {"rows": len(results), "valid": valid, "errors": len(results) - valid,
            "created": created, "committed": commit, "results": results}


def template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Envíos"
    cols = ["origen", "destino", "carga", "cliente", "peso_t", "salida", *OPTIONAL]
    ws.append(cols)
    ws.append(["Bogotá", "Cartagena", "Café excelso en sacos", "Cafés de la Cordillera", 25,
               datetime.now().replace(minute=0, second=0, microsecond=0), "balanceado", 0, "", ""])
    ws.append(["Barrancabermeja", "Cartagena", "Combustibles líquidos", "Caribe Energía", 900,
               datetime.now().replace(minute=0, second=0, microsecond=0), "costo", 0, "fluvial", ""])
    bold, fill = Font(bold=True), PatternFill("solid", fgColor="EEF1F6")
    for c in range(1, len(cols) + 1):
        ws.cell(row=1, column=c).font = bold
        ws.cell(row=1, column=c).fill = fill
        ws.column_dimensions[get_column_letter(c)].width = 22
    for r in (2, 3):
        ws.cell(row=r, column=6).number_format = "yyyy-mm-dd hh:mm"
    ws.freeze_panes = "A2"
    notes = wb.create_sheet("Instrucciones")
    for line in [
        ["Columna", "Qué va"],
        ["origen / destino", "Ciudad tal como aparece en el simulador (por ejemplo, Bogotá, Leticia)."],
        ["carga / cliente", "Texto libre."],
        ["peso_t", "Toneladas, entre 0 y 20.000."],
        ["salida", "Fecha y hora de salida: AAAA-MM-DD HH:MM o DD/MM/AAAA HH:MM."],
        ["prioridad", "Opcional: tiempo, costo o balanceado (por defecto)."],
        ["retraso_h", "Opcional: horas de retraso operativo."],
        ["modos", "Opcional: obliga a usar solo estos modos, separados por coma (terrestre, fluvial, maritimo, aereo, ferreo)."],
        ["vias", "Opcional: vías por las que debe pasar, separadas por punto y coma."],
        ["", "Máximo 500 filas por carga. Primero se muestra una vista previa; nada se guarda hasta confirmar."],
    ]:
        notes.append(line)
    notes.column_dimensions["A"].width = 20
    notes.column_dimensions["B"].width = 100
    notes["A1"].font = bold
    notes["B1"].font = bold
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
