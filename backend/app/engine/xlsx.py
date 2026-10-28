"""
xlsx.py
Exportación de envíos a Excel (openpyxl). Una hoja "Envíos" con una fila por envío y una hoja "Notas"
con la moneda, la TRM usada y la fecha de generación. Los montos se guardan como números (no como
texto) para que el planificador pueda sumar y filtrar en Excel.
"""
from __future__ import annotations

import io
import json
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.api import format as F
from app.engine.dispatch import status_at

HEADER_FILL = PatternFill("solid", fgColor="F5F5F7")


def render(rows: list, currency: str, fx: dict, now: datetime) -> bytes:
    rate = fx.get("rate") if currency == "COP" else None
    if currency == "COP" and not rate:
        raise ValueError("No hay TRM configurada: no se puede exportar en COP.")
    factor = rate or 1.0
    money_fmt = '#,##0' if currency == "COP" else '#,##0.00'

    wb = Workbook()
    ws = wb.active
    ws.title = "Envíos"
    cols = [
        ("Código", 14), ("Origen", 18), ("Destino", 18), ("Carga", 30), ("Cliente", 24), ("Estado", 18),
        ("Modos", 26), ("Transbordos", 12), ("Peso (t)", 10), ("Distancia (km)", 14), ("Tiempo (h)", 11),
        (f"Flete ({currency})", 14), ("Salida", 17), ("ETA", 17), ("Retraso (h)", 11), ("Prioridad", 12),
        ("Estado de la ruta", 16), ("Nota de la ruta", 40),
    ]
    for i, (name, width) in enumerate(cols, start=1):
        c = ws.cell(row=1, column=i, value=name)
        c.font = Font(bold=True)
        c.fill = HEADER_FILL
        c.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"

    for r in rows:
        modes = json.loads(r["modes"]) if isinstance(r["modes"], str) else (r["modes"] or [])
        status = status_at(r, now) if callable(status_at) else r["status"]
        ws.append([
            r["shipment_code"], r["origin_city"], r["dest_city"], r["cargo"], r["client"],
            status if isinstance(status, str) else r["status"],
            " → ".join(F.MODE_LABELS.get(m, m) for m in modes), r["n_transfers"],
            float(r["weight_t"]), float(r["distance_km"] or 0), float(r["total_time_h"] or 0),
            round(float(r["total_cost"] or 0) * factor, 2),
            datetime.fromisoformat(r["departure_at"]), datetime.fromisoformat(r["eta_at"]),
            float(r["delay_h"] or 0), r["priority"].capitalize(),
            "Sin ruta" if r.get("route_status") == "sin_ruta" else "Con ruta", r.get("route_note") or "",
        ])
    for row in ws.iter_rows(min_row=2):
        row[8].number_format = '#,##0.0'
        row[9].number_format = '#,##0'
        row[10].number_format = '#,##0.0'
        row[11].number_format = money_fmt
        row[12].number_format = "dd/mm/yyyy hh:mm"
        row[13].number_format = "dd/mm/yyyy hh:mm"
        row[14].number_format = '#,##0.0'
    ws.auto_filter.ref = ws.dimensions

    notes = wb.create_sheet("Notas")
    notes["A1"], notes["B1"] = "Generado", now.strftime("%d/%m/%Y %H:%M")
    notes["A2"], notes["B2"] = "Moneda", currency
    if rate:
        notes["A3"], notes["B3"] = "TRM (COP por USD)", rate
        notes["A4"], notes["B4"] = "Fecha y fuente de la TRM", f'{fx.get("date") or "—"} · {fx.get("source") or "—"}'
    notes["A6"] = "Westvane TMS. Datos simulados: clientes, transportistas y telemetría son ficticios."
    for c in ("A1", "A2", "A3", "A4"):
        notes[c].font = Font(bold=True)
    notes.column_dimensions["A"].width = 26
    notes.column_dimensions["B"].width = 40

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
